"""Professional Styled Excel Audit Workbook Generator."""

from pathlib import Path
from decimal import Decimal
import openpyxl  # type: ignore
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore

from app.core.exceptions import ExportError
from app.domain.enums import MatchStatus
from app.domain.models import MatchRecord, ReconciliationSummary

# Color Palette for Accounting Spreadsheet
COLOR_HEADER_BG = "1E293B"     # Dark Slate
COLOR_HEADER_FG = "FFFFFF"     # White
COLOR_MATCHED_BG = "D1FAE5"    # Soft Emerald
COLOR_MATCHED_FG = "065F46"
COLOR_DIFF_BG = "FEF3C7"       # Soft Amber
COLOR_DIFF_FG = "92400E"
COLOR_MISSING_A_BG = "FEE2E2"  # Soft Red
COLOR_MISSING_A_FG = "991B1B"
COLOR_MISSING_B_BG = "EDE9FE"  # Soft Purple
COLOR_MISSING_B_FG = "5B21B6"

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _apply_header_style(ws, row_idx: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = Font(name="Segoe UI", size=11, bold=True, color=COLOR_HEADER_FG)
        cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autofit_columns(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)


class ExcelExporter:
    """Exports reconciliation results into an audit-ready formatted Excel workbook."""

    @classmethod
    def export(cls, summary: ReconciliationSummary, output_path: Path) -> Path:
        try:
            wb = openpyxl.Workbook()
            # Remove default active sheet
            default_sheet = wb.active

            # -------------------------------------------------------------
            # Sheet 1: Executive Summary
            # -------------------------------------------------------------
            ws_sum = wb.create_sheet(title="Executive Summary")
            cls._build_summary_sheet(ws_sum, summary)

            # -------------------------------------------------------------
            # Sheet 2: Mismatched Records
            # -------------------------------------------------------------
            ws_mismatch = wb.create_sheet(title="Tax Differences")
            cls._build_match_table_sheet(
                ws_mismatch,
                [m for m in summary.matches if m.match_status == MatchStatus.MATCHED_WITH_DIFFERENCE],
                title="Matched Invoices with Differences (Exceeding Tolerance)",
            )

            # -------------------------------------------------------------
            # Sheet 3: Missing in GSTR-2B (ITC at Risk)
            # -------------------------------------------------------------
            ws_risk = wb.create_sheet(title="ITC at Risk (Missing in 2B)")
            cls._build_missing_sheet(
                ws_risk,
                [m for m in summary.matches if m.match_status == MatchStatus.MISSING_IN_SOURCE_A],
                is_itc_at_risk=True,
            )

            # -------------------------------------------------------------
            # Sheet 4: Missing in Books (Unclaimed ITC)
            # -------------------------------------------------------------
            ws_unclaimed = wb.create_sheet(title="Unclaimed (Missing in Books)")
            cls._build_missing_sheet(
                ws_unclaimed,
                [m for m in summary.matches if m.match_status == MatchStatus.MISSING_IN_SOURCE_B],
                is_itc_at_risk=False,
            )

            # -------------------------------------------------------------
            # Sheet 5: Matched Records
            # -------------------------------------------------------------
            ws_matched = wb.create_sheet(title="Matched Invoices")
            cls._build_match_table_sheet(
                ws_matched,
                [m for m in summary.matches if m.match_status == MatchStatus.MATCHED],
                title="Perfect Matches (Within Tolerance)",
            )

            # Clean up default sheet
            if default_sheet:
                wb.remove(default_sheet)

            output_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            return output_path
        except Exception as e:
            raise ExportError(f"Failed to generate Excel report: {e}")

    @classmethod
    def _build_summary_sheet(cls, ws, summary: ReconciliationSummary) -> None:
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = "GST RECONCILIATION AUDIT REPORT"
        title_cell.font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Metadata
        ws["A3"] = "Project / Entity:"
        ws["B3"] = summary.project_name
        ws["A4"] = "Reconciled On:"
        ws["B4"] = summary.created_at
        ws["A5"] = "Source A (Portal):"
        ws["B5"] = summary.source_a_file or "N/A"
        ws["A6"] = "Source B (Books):"
        ws["B6"] = summary.source_b_file or "N/A"

        for r in range(3, 7):
            ws[f"A{r}"].font = Font(name="Segoe UI", size=10, bold=True, color="475569")
            ws[f"B{r}"].font = Font(name="Segoe UI", size=10, color="1E293B")

        # KPI Metrics Table
        headers = ["Reconciliation Metric", "Invoice Count", "Percentage", "Financial Impact"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=8, column=c_idx, value=h)
        _apply_header_style(ws, 8, len(headers))
        ws.row_dimensions[8].height = 24

        rows_data = [
            ("Matched (Within Tolerance)", summary.total_matched, f"{summary.match_rate_percentage}%", f"₹ {summary.total_taxable_a:,.2f}"),
            ("Matched with Differences", summary.total_matched_with_diff, f"{(summary.total_matched_with_diff / max(summary.total_processed, 1) * 100):.1f}%", f"₹ {summary.net_diff_tax:,.2f} Delta"),
            ("ITC at Risk (Missing in GSTR-2B)", summary.total_missing_in_a, f"{(summary.total_missing_in_a / max(summary.total_processed, 1) * 100):.1f}%", f"₹ {summary.itc_at_risk_amount:,.2f}"),
            ("Unclaimed ITC (Missing in Books)", summary.total_missing_in_b, f"{(summary.total_missing_in_b / max(summary.total_processed, 1) * 100):.1f}%", f"₹ {summary.unclaimed_itc_amount:,.2f}"),
            ("Duplicates in Source A / B", summary.total_duplicates_a + summary.total_duplicates_b, "N/A", "Review Required"),
            ("Total Records Processed", summary.total_processed, "100.0%", f"Portal: {summary.total_records_a} | Books: {summary.total_records_b}"),
        ]

        for r_idx, row in enumerate(rows_data, 9):
            ws.cell(row=r_idx, column=1, value=row[0]).font = Font(name="Segoe UI", size=10, bold=(r_idx == 14))
            ws.cell(row=r_idx, column=2, value=row[1]).alignment = Alignment(horizontal="right")
            ws.cell(row=r_idx, column=3, value=row[2]).alignment = Alignment(horizontal="right")
            ws.cell(row=r_idx, column=4, value=row[3]).alignment = Alignment(horizontal="right")
            for c in range(1, 5):
                ws.cell(row=r_idx, column=c).border = BORDER_THIN

        _autofit_columns(ws)

    @classmethod
    def _build_match_table_sheet(cls, ws, matches: list[MatchRecord], title: str) -> None:
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Supplier GSTIN",
            "Supplier Name",
            "Portal Inv No",
            "Books Inv No",
            "Invoice Date",
            "Match Level",
            "Taxable (Portal)",
            "Taxable (Books)",
            "Diff Taxable",
            "Tax (Portal)",
            "Tax (Books)",
            "Diff Total Tax",
            "Review Status",
            "Diagnosis / Explanation",
        ]

        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=h)
        _apply_header_style(ws, 1, len(headers))
        ws.row_dimensions[1].height = 26

        for r_idx, m in enumerate(matches, 2):
            ra = m.record_a
            rb = m.record_b

            gstin = (rb.supplier_gstin if rb else (ra.supplier_gstin if ra else ""))
            name = (rb.supplier_name if rb and rb.supplier_name else (ra.supplier_name if ra else ""))
            inv_a = ra.raw_invoice_number if ra else ""
            inv_b = rb.raw_invoice_number if rb else ""
            rec_for_date = rb if (rb and rb.invoice_date) else ra
            inv_date = rec_for_date.invoice_date.isoformat() if (rec_for_date and rec_for_date.invoice_date) else ""

            taxable_a = float(ra.taxable_value) if ra else 0.0
            taxable_b = float(rb.taxable_value) if rb else 0.0
            tax_a = float(ra.calculate_total_tax()) if ra else 0.0
            tax_b = float(rb.calculate_total_tax()) if rb else 0.0

            row_vals = [
                gstin,
                name,
                inv_a,
                inv_b,
                inv_date,
                m.match_level.value,
                taxable_a,
                taxable_b,
                float(m.diff_taxable),
                tax_a,
                tax_b,
                float(m.diff_total_tax),
                m.review_status.value,
                m.explanation,
            ]

            for c_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = BORDER_THIN
                if c_idx in (7, 8, 9, 10, 11, 12):
                    cell.number_format = "₹ #,##0.00"
                    cell.alignment = Alignment(horizontal="right")

        _autofit_columns(ws)

    @classmethod
    def _build_missing_sheet(cls, ws, matches: list[MatchRecord], is_itc_at_risk: bool) -> None:
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Supplier GSTIN",
            "Supplier Name",
            "Invoice Number",
            "Invoice Date",
            "Taxable Value",
            "IGST",
            "CGST",
            "SGST",
            "Cess",
            "Total Tax",
            "Statutory Action Required",
        ]

        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=h)
        _apply_header_style(ws, 1, len(headers))
        ws.row_dimensions[1].height = 26

        for r_idx, m in enumerate(matches, 2):
            rec = m.record_b if is_itc_at_risk else m.record_a
            if not rec:
                continue

            action = (
                "Follow up with supplier to file GSTR-1"
                if is_itc_at_risk
                else "Book invoice in ERP to claim eligible ITC"
            )

            row_vals = [
                rec.supplier_gstin,
                rec.supplier_name or "",
                rec.raw_invoice_number,
                rec.invoice_date.isoformat() if rec.invoice_date else "",
                float(rec.taxable_value),
                float(rec.igst),
                float(rec.cgst),
                float(rec.sgst),
                float(rec.cess),
                float(rec.calculate_total_tax()),
                action,
            ]

            for c_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = BORDER_THIN
                if c_idx in (5, 6, 7, 8, 9, 10):
                    cell.number_format = "₹ #,##0.00"
                    cell.alignment = Alignment(horizontal="right")

        _autofit_columns(ws)
